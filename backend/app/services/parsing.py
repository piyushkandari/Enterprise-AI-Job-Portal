import io
from fastapi import HTTPException, status

MAX_BYTES = 5 * 1024 * 1024  # 5 MB — plenty for any resume
MIN_TEXT_CHARS = 100          # below this, the PDF was probably a scan/image


def extract_text(filename: str, data: bytes) -> str:
    lower = filename.lower()

    if lower.endswith(".pdf"):
        text = _extract_pdf(data)
    elif lower.endswith(".docx"):
        text = _extract_docx(data)
    elif lower.endswith(".txt"):
        text = _extract_txt(data)
    else:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            "unsupported file type — upload PDF, DOCX, or TXT",
        )

    text = text.strip()
    if len(text) < MIN_TEXT_CHARS:
        raise HTTPException(
            status.HTTP_422_UNPROCESSABLE_ENTITY,
            "Could not extract readable text. If your resume is a scanned image, "
            "please upload a text-based PDF or DOCX instead.",
        )
    return text


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader
    try:
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception as e:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY,
                            f"Failed to read PDF: {e}")


def _extract_docx(data: bytes) -> str:
    from docx import Document
    try:
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        # also pull text out of tables — resumes often use them for layout
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return "\n".join(parts)
    except Exception as e:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY,
                            f"Failed to read DOCX: {e}")


def _extract_txt(data: bytes) -> str:
    try:
        return data.decode("utf-8", errors="replace")
    except Exception as e:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY,
                            f"Failed to read TXT: {e}")